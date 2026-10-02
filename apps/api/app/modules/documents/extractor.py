from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re


class ExtractionError(RuntimeError):
    pass


class EncryptedDocumentError(ExtractionError):
    pass


@dataclass(frozen=True, slots=True)
class ExtractedChunk:
    ordinal: int
    locator_type: str
    locator: str
    text: str


def _clean(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[\t ]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def extract_document(path: Path, extension: str) -> list[ExtractedChunk]:
    if extension == ".pdf":
        return _pdf(path)
    if extension == ".pptx":
        return _pptx(path)
    if extension == ".docx":
        return _docx(path)
    if extension == ".txt":
        return _txt(path)
    raise ExtractionError("No extractor is available for this file type.")


def _pdf(path: Path) -> list[ExtractedChunk]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            try:
                result = reader.decrypt("")
            except Exception as exc:
                raise EncryptedDocumentError("Encrypted PDFs are not supported. Remove the password and upload again.") from exc
            if not result:
                raise EncryptedDocumentError("Encrypted PDFs are not supported. Remove the password and upload again.")
        chunks: list[ExtractedChunk] = []
        for index, page in enumerate(reader.pages, start=1):
            text = _clean(page.extract_text() or "")
            if text:
                chunks.append(ExtractedChunk(index, "page", f"Page {index}", text))
    except EncryptedDocumentError:
        raise
    except Exception as exc:
        raise ExtractionError("The PDF is corrupt or could not be parsed safely.") from exc
    if not chunks:
        raise ExtractionError("No extractable PDF text was found. OCR for scanned/handwritten pages is outside this story.")
    return chunks


def _pptx(path: Path) -> list[ExtractedChunk]:
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
        chunks: list[ExtractedChunk] = []
        for index, slide in enumerate(presentation.slides, start=1):
            pieces: list[str] = []
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text and text.strip():
                    pieces.append(text)
            clean = _clean("\n".join(pieces))
            if clean:
                chunks.append(ExtractedChunk(index, "slide", f"Slide {index}", clean))
    except Exception as exc:
        raise ExtractionError("The PowerPoint file is corrupt or could not be parsed safely.") from exc
    if not chunks:
        raise ExtractionError("No extractable slide text was found.")
    return chunks


def _docx(path: Path) -> list[ExtractedChunk]:
    from docx import Document

    try:
        document = Document(str(path))
        paragraphs = [_clean(item.text) for item in document.paragraphs if _clean(item.text)]
    except Exception as exc:
        raise ExtractionError("The Word document is corrupt or could not be parsed safely.") from exc
    if not paragraphs:
        raise ExtractionError("No extractable text was found in the document.")
    chunks: list[ExtractedChunk] = []
    for start in range(0, len(paragraphs), 12):
        end = min(start + 12, len(paragraphs))
        chunks.append(ExtractedChunk(len(chunks) + 1, "paragraphs", f"Paragraphs {start + 1}–{end}", "\n".join(paragraphs[start:end])))
    return chunks


def _txt(path: Path) -> list[ExtractedChunk]:
    text = _clean(path.read_text(encoding="utf-8"))
    if not text:
        raise ExtractionError("The text file does not contain usable text.")
    chunks: list[ExtractedChunk] = []
    parts = [text[index:index + 3500] for index in range(0, len(text), 3500)]
    for index, part in enumerate(parts, start=1):
        chunks.append(ExtractedChunk(index, "section", f"Section {index}", part))
    return chunks
